"""Daily Excel report."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from job_copilot.config import IST, REPORTS_DIR, ensure_dirs

COLUMNS = [
    ("Portal", "portal", 10), ("Job Title", "title", 34), ("Company", "company", 26),
    ("Employer Email", "recruiter_email", 26), ("Employer Phone", "recruiter_phone", 16),
    ("Match Score", "match_score", 11), ("Resume URL", "resume_url", 40), ("Status", "status", 15),
    ("Why", "why", 50), ("Job URL", "url", 40), ("Posted", "posted_at", 18),
]


def build_report(jobs: list[dict], day: datetime | None = None) -> Path:
    ensure_dirs()
    day = day or datetime.now(IST)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([c[0] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A5FB4")
    for j in jobs:
        why = (j.get("match") or {}).get("explanation") or j.get("status_reason") or ""
        row = {**j, "why": why}
        ws.append([row.get(key) or "" for _, key, _ in COLUMNS])
        for idx in (7, 10):  # make URLs clickable
            cell = ws.cell(row=ws.max_row, column=idx)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="1A5FB4", underline="single")
    for i, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary.append(["Status", "Count"])
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    for status, n in Counter(j["status"] for j in jobs).most_common():
        summary.append([status, n])

    path = REPORTS_DIR / f"job_report_{day:%Y-%m-%d}.xlsx"
    wb.save(path)
    return path


def summary_text(jobs: list[dict]) -> str:
    counts = Counter(j["status"] for j in jobs)
    lines = [f"Jobs processed: {len(jobs)}"] + [f"  {k}: {v}" for k, v in counts.most_common()]
    top = [j for j in jobs if j["status"] in ("pending_review", "approved", "applied")][:10]
    if top:
        lines += ["", "Top matches:"]
        lines += [f"  [{j['match_score']}] {j['title']} @ {j['company']} - {j['status']}" for j in top]
    if counts.get("pending_review"):
        lines += ["", "Review pending jobs with:  python main.py review"]
    return "\n".join(lines)
